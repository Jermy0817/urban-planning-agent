"""Local, evidence-first planning workflow. No network or model calls."""
import re
from pathlib import Path
from dataclasses import dataclass, asdict

@dataclass
class Passage:
    source: str
    location: str
    text: str

def read_document(path):
    p = Path(path)
    if p.stat().st_size > 300 * 1024 * 1024:
        raise ValueError('文件超过 300 MB，请拆分后导入')
    result = []
    def add(loc, text):
        if text and str(text).strip():
            result.append(Passage(p.name, loc, str(text).strip()))
    ext = p.suffix.lower()
    if ext in ('.txt', '.md'):
        for i, line in enumerate(p.read_text(encoding='utf-8-sig').splitlines(), 1):
            add(f'第 {i} 行', line)
    elif ext == '.pdf':
        from pypdf import PdfReader
        for i, page in enumerate(PdfReader(p).pages, 1):
            text = page.extract_text() or ''
            add(f'第 {i} 页', text or '[未提取到文字：可能需要 OCR，不能视为无内容]')
    elif ext == '.docx':
        from docx import Document
        doc = Document(p)
        for i, para in enumerate(doc.paragraphs, 1):
            add(f'正文段落 {i}', para.text)
        for t, table in enumerate(doc.tables, 1):
            for r, row in enumerate(table.rows, 1):
                add(f'表 {t} 行 {r}', ' | '.join(c.text for c in row.cells))
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(p, read_only=True, data_only=False, keep_links=False)
        try:
            for sheet in book:
                for row in sheet.iter_rows():
                    cells = [f'{c.coordinate}={c.value}' for c in row if c.value is not None]
                    if cells:
                        add(f'{sheet.title} 第 {row[0].row} 行', ' | '.join(cells))
        finally:
            book.close()
    else:
        raise ValueError('仅支持 PDF、DOCX、XLSX、UTF-8 TXT/MD；旧格式请另存为新格式')
    if not result:
        raise ValueError('没有提取到文本，请检查文件或先做 OCR')
    return result

def facts(passages):
    patterns = {'面积': r'\d[\d,.]*\s*(?:万平方米|平方米|万㎡|㎡|公顷|亩|m²)',
                '投资': r'\d[\d,.]*\s*(?:亿元|万元|元)',
                '建设内容': r'[^。\n]*(?:建设内容|拟建|新建|改造|扩建|建设规模)[^。\n]*'}
    rows = []
    for p in passages:
        for category, pattern in patterns.items():
            for m in re.finditer(pattern, p.text):
                rows.append(dict(category=category, value=m.group(), source=p.source,
                                 location=p.location, evidence=p.text[max(0,m.start()-60):m.end()+100],
                                 status='待人工核对'))
    return rows

def retrieve(query, passages, limit=5):
    clean = re.sub(r'\s+', '', query.lower())
    terms = set(clean[i:i+2] for i in range(max(0,len(clean)-1)))
    if len(clean) == 1:
        terms = {clean}
    ranked = [(sum(t in p.text.lower() for t in terms), p) for p in passages]
    return [p for score,p in sorted(ranked, key=lambda x:x[0], reverse=True) if score][:limit]

def report(project, experience, query):
    rows = facts(project)
    lines = ['# 策划初稿（规则生成，未经审核）', '', '不构成政策适用性、投资可行性或审批结论。', '', '## 项目事实候选']
    for row in rows:
        lines += [f"- {row['category']}：{row['value']}【{row['source']} / {row['location']}】", f"  原文：{row['evidence']}"]
    if not rows:
        lines += ['未识别到事实；不代表原文不存在，请人工补录。']
    lines += ['', '## 经验参考（关键词匹配，不代表适用）']
    hits = retrieve(query, experience)
    for p in hits:
        lines += [f'- 【{p.source} / {p.location}】{p.text}', '  建议：核对该经验的地区、年份、项目类型和实施条件后，再决定是否借鉴。']
    if not hits:
        lines += ['暂无匹配经验，请补充既有项目复盘。']
    lines += ['', '## 策划建议与下一步', '- 明确服务人群、需求证据、运营主体与收入/支出结构，形成至少两个规模方案。']
    for category, advice in [('面积','分别核对用地、建筑、计容面积及单位，不能直接合并不同口径。'), ('投资','分别核对总投资、建安费、设备费与运营费用；明确价格年份和资金来源。'), ('建设内容','将建设内容拆成必要、可选、分期实施三类，逐项说明需求依据。')]:
        relevant = [r for r in rows if r['category']==category]
        refs = '；'.join(dict.fromkeys(f"{r['source']} / {r['location']}" for r in relevant))
        lines += [f'- {advice} 依据：{refs or "资料缺失，需补充"}']
    lines += ['', '## 人工审核清单', '- [ ] 原文与数值、单位、口径逐项核对；相互冲突的数据不自行择一。', '- [ ] 政策核对发布机关、文号、地区、有效期及废止/修订情况。', '- [ ] 案例与公开数据记录原始网址、发布日期、访问日期和统计口径。', '- [ ] 经验适用性及建议由项目负责人审核；当前草稿不得直接用于报批。']
    return '\n'.join(lines)
