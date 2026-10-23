import tempfile
import unittest
from pathlib import Path
from core import Passage, facts, retrieve, read_document, report

class WorkflowTests(unittest.TestCase):
    def test_evidence_and_conflicts(self):
        p=[Passage('a','页1','面积12亩，总投资3000万元。'),Passage('b','页2','面积15亩。')]
        rows=facts(p)
        self.assertEqual([r['value'] for r in rows if r['category']=='面积'],['12亩','15亩'])
        self.assertEqual(rows[0]['location'],'页1')
        self.assertIn('3000万元',report(p,[],''))

    def test_no_irrelevant_experience(self):
        self.assertEqual(retrieve('医院',[Passage('a','1','社区运营')]),[])

    def test_readers(self):
        from docx import Document
        from openpyxl import Workbook
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory() as d:
            root=Path(d)
            doc=Document();doc.add_paragraph('面积12亩')
            doc.add_table(rows=1,cols=1).cell(0,0).text='总投资3000万元'
            doc.save(root/'a.docx')
            self.assertEqual(len(facts(read_document(root/'a.docx'))),2)
            wb=Workbook();wb.active['A1']='建筑面积';wb.active['B1']='8000平方米'
            wb.active['C1']='=1+1';wb.save(root/'a.xlsx')
            text=read_document(root/'a.xlsx')[0].text
            self.assertIn('B1=8000平方米',text);self.assertIn('=1+1',text)
            pdf=PdfWriter();pdf.add_blank_page(width=100,height=100);pdf.write(root/'a.pdf')
            self.assertIn('OCR',read_document(root/'a.pdf')[0].text)

if __name__=='__main__':
    unittest.main()
